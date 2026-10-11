"""
Unified Multi-Format Document Ingestion Engine for SpecGuard.
Extracts normalized DocumentModel with high-fidelity coordinate, font, and table preservation.
Supports PDF, DOCX, XLSX, TXT, and Images (PNG, JPG, JPEG, TIFF).
"""

import os
import re
import hashlib
from pathlib import Path
from typing import List, Dict, Optional, Any
import logging

import pymupdf  # PyMuPDF
import cv2
import numpy as np

from specguard.core.models import (
    DocumentModel, PageModel, TextBlock, TableData, FigureData, BBox,
    WordModel, LineModel
)
from specguard.core.scanned_pipeline import ScannedDocumentPipeline
from specguard.core.cv_layout import LayoutAnalyzerCV
from specguard.core.reading_order import ReadingOrderEngine
from specguard.core.docx_extractor import DocxExtractor

logger = logging.getLogger(__name__)


class DocumentParser:
    """Dispatches document parsing to specialized format handlers."""

    @staticmethod
    def parse_file(
        file_path: str,
        progress_callback=None,
        cancellation_token=None,
        deep_analysis: bool = False
    ) -> DocumentModel:
        path = Path(file_path).resolve()
        if not path.exists():
            raise FileNotFoundError(f"Document not found: {file_path}")

        file_size = path.stat().st_size
        sha256 = hashlib.sha256()
        with open(path, "rb") as f:
            while chunk := f.read(65536):
                sha256.update(chunk)
        file_hash = sha256.hexdigest()

        ext = path.suffix.lower()
        if ext == ".pdf":
            return DocumentParser._parse_pdf(
                path, file_hash, file_size,
                progress_callback=progress_callback,
                cancellation_token=cancellation_token,
                deep_analysis=deep_analysis
            )
        elif ext == ".docx":
            if progress_callback:
                try:
                    progress_callback(1, 1, "Parsing Word document XML structure...")
                except Exception:
                    pass
            return DocxExtractor.extract(path, file_hash, file_size)
        elif ext in [".xlsx", ".xls"]:
            if progress_callback:
                try:
                    progress_callback(1, 1, "Parsing Excel spreadsheets & tabular ranges...")
                except Exception:
                    pass
            return DocumentParser._parse_xlsx(path, file_hash, file_size)
        elif ext == ".txt":
            if progress_callback:
                try:
                    progress_callback(1, 1, "Reading plain text specification...")
                except Exception:
                    pass
            return DocumentParser._parse_txt(path, file_hash, file_size)
        elif ext in [".png", ".jpg", ".jpeg", ".tiff", ".tif"]:
            if progress_callback:
                try:
                    progress_callback(1, 1, "Preprocessing image & computer vision layout...")
                except Exception:
                    pass
            return DocumentParser._parse_image(path, file_hash, file_size)
        else:
            raise ValueError(f"Unsupported document format: {ext}")

    @staticmethod
    def _parse_pdf(
        path: Path,
        file_hash: str,
        file_size: int,
        progress_callback=None,
        cancellation_token=None,
        deep_analysis: bool = False
    ) -> DocumentModel:
        doc = pymupdf.open(str(path))
        total_pages = len(doc)
        pages: List[PageModel] = []
        reading_engine = ReadingOrderEngine()

        for page_idx, fitz_page in enumerate(doc):
            if cancellation_token and getattr(cancellation_token, "is_cancelled", False):
                doc.close()
                raise InterruptedError("Document processing was cancelled by user.")

            page_num = page_idx + 1
            if progress_callback:
                try:
                    progress_callback(page_num, total_pages, f"Extracting page {page_num} of {total_pages}")
                except Exception:
                    pass

            rect = fitz_page.rect
            page_w = float(rect.width)
            page_h = float(rect.height)
            rotation = fitz_page.rotation

            raw_blocks: List[TextBlock] = []
            tables: List[TableData] = []
            figures: List[FigureData] = []

            # Extract detailed text blocks with font and coordinate data
            page_dict = fitz_page.get_text("dict")

            # Extract exact word coordinates from PyMuPDF
            words_by_line: Dict[Tuple[int, int], List[Tuple[float, float, float, float, str]]] = {}
            try:
                for w_item in fitz_page.get_text("words"):
                    # w_item: (x0, y0, x1, y1, word_text, block_no, line_no, word_no)
                    k = (int(w_item[5]), int(w_item[6]))
                    if k not in words_by_line:
                        words_by_line[k] = []
                    words_by_line[k].append((float(w_item[0]), float(w_item[1]), float(w_item[2]), float(w_item[3]), str(w_item[4])))
            except Exception as w_err:
                logger.debug("Native word extraction error on page %d: %s", page_num, w_err)
                words_by_line = {}

            block_id = 0
            for b_idx, b in enumerate(page_dict.get("blocks", [])):
                if b.get("type") == 0:  # Text block
                    block_words: List[WordModel] = []
                    block_lines: List[LineModel] = []

                    for l_idx, line in enumerate(b.get("lines", [])):
                        line_text_parts = []
                        line_bbox = line.get("bbox", (0, 0, 0, 0))
                        font_name = "Helvetica"
                        font_size = 11.0
                        line_words: List[WordModel] = []

                        # Gather span styles for this line
                        spans = line.get("spans", [])
                        span_info = []
                        for span in spans:
                            span_text = span.get("text", "")
                            line_text_parts.append(span_text)
                            s_font = span.get("font", font_name)
                            s_size = float(span.get("size", font_size))
                            flags = span.get("flags", 0)
                            s_italic = bool(flags & 2)
                            s_bold = bool(flags & 16 or "bold" in s_font.lower())
                            s_bbox = span.get("bbox", (0, 0, 0, 0))
                            span_info.append({
                                "text": span_text,
                                "font": s_font,
                                "size": s_size,
                                "bold": s_bold,
                                "italic": s_italic,
                                "bbox": s_bbox
                            })
                            font_name = s_font
                            font_size = s_size

                        # Get exact words for this block and line if available
                        exact_words = words_by_line.get((b_idx, l_idx), [])
                        if exact_words:
                            for wx0, wy0, wx1, wy1, wtext in exact_words:
                                # Determine matching span for font attributes
                                matched_span = None
                                w_mid_x = (wx0 + wx1) / 2.0
                                for si in span_info:
                                    sb = si["bbox"]
                                    if sb[0] - 2.0 <= w_mid_x <= sb[2] + 2.0:
                                        matched_span = si
                                        break
                                if not matched_span and span_info:
                                    matched_span = span_info[0]

                                w_model = WordModel(
                                    text=wtext,
                                    bbox=BBox(x0=wx0, y0=wy0, x1=wx1, y1=wy1),
                                    font_name=matched_span["font"] if matched_span else font_name,
                                    font_size=round(matched_span["size"] if matched_span else font_size, 1),
                                    is_bold=matched_span["bold"] if matched_span else False,
                                    is_italic=matched_span["italic"] if matched_span else False
                                )
                                line_words.append(w_model)
                                block_words.append(w_model)
                        else:
                            # Fallback: Proportional character slicing inside each span
                            for si in span_info:
                                span_text = si["text"]
                                sb = si["bbox"]
                                s_w = sb[2] - sb[0]
                                s_len = max(1, len(span_text))
                                char_w = s_w / s_len
                                c_offset = 0
                                for word_str in span_text.split():
                                    if not word_str.strip():
                                        continue
                                    idx_in_span = span_text.find(word_str, c_offset)
                                    if idx_in_span == -1:
                                        idx_in_span = c_offset
                                    w_x0 = sb[0] + idx_in_span * char_w
                                    w_x1 = w_x0 + len(word_str) * char_w
                                    c_offset = idx_in_span + len(word_str)

                                    w_model = WordModel(
                                        text=word_str,
                                        bbox=BBox(x0=w_x0, y0=sb[1], x1=w_x1, y1=sb[3]),
                                        font_name=si["font"],
                                        font_size=round(si["size"], 1),
                                        is_bold=si["bold"],
                                        is_italic=si["italic"]
                                    )
                                    line_words.append(w_model)
                                    block_words.append(w_model)

                        full_line_text = "".join(line_text_parts).strip()
                        if full_line_text:
                            block_lines.append(LineModel(
                                words=line_words,
                                text=full_line_text,
                                bbox=BBox(x0=line_bbox[0], y0=line_bbox[1], x1=line_bbox[2], y1=line_bbox[3]),
                                font_size=round(font_size, 1)
                            ))

                    if block_lines:
                        # Compute overall block bbox from lines and clamp to page dimensions
                        bx0 = max(0.0, min(page_w, min(l.bbox.x0 for l in block_lines if l.bbox)))
                        by0 = max(0.0, min(page_h, min(l.bbox.y0 for l in block_lines if l.bbox)))
                        bx1 = max(0.0, min(page_w, max(l.bbox.x1 for l in block_lines if l.bbox)))
                        by1 = max(0.0, min(page_h, max(l.bbox.y1 for l in block_lines if l.bbox)))
                        b_text = "\n".join(l.text for l in block_lines)

                        raw_blocks.append(TextBlock(
                            text=b_text,
                            bbox=BBox(x0=bx0, y0=by0, x1=bx1, y1=by1),
                            font_name=block_lines[0].words[0].font_name if block_words else "Helvetica",
                            font_size=block_lines[0].font_size,
                            is_bold=block_words[0].is_bold if block_words else False,
                            is_italic=block_words[0].is_italic if block_words else False,
                            block_id=block_id,
                            words=block_words,
                            lines=block_lines
                        ))
                        block_id += 1

            total_native_chars = sum(len(b.text.strip()) for b in raw_blocks)
            page_area = max(1.0, page_w * page_h)
            is_scanned_page = False

            # Inspect embedded raster image objects independently of rendered pages
            for b in page_dict.get("blocks", []):
                if b.get("type") == 1:
                    img_bbox = b.get("bbox", (0, 0, 0, 0))
                    ix0 = max(0.0, min(page_w, float(img_bbox[0])))
                    iy0 = max(0.0, min(page_h, float(img_bbox[1])))
                    ix1 = max(0.0, min(page_w, float(img_bbox[2])))
                    iy1 = max(0.0, min(page_h, float(img_bbox[3])))
                    iw = max(0.0, ix1 - ix0)
                    ih = max(0.0, iy1 - iy0)
                    if iw < 10.0 or ih < 10.0:
                        continue

                    i_area = iw * ih
                    coverage = i_area / page_area

                    # Prevent page-background contours and scan boundaries from being reported as ordinary figures
                    if coverage >= 0.85:
                        if total_native_chars < 50:
                            is_scanned_page = True
                        # Do not report full-page background or scan layer as an ordinary figure
                        continue
                    elif coverage >= 0.70 and (ix0 <= 25 and iy0 <= 25 and ix1 >= page_w - 25 and iy1 >= page_h - 25):
                        if total_native_chars < 50:
                            is_scanned_page = True
                        continue
                    else:
                        # Genuine embedded raster image/diagram
                        if iw >= 20.0 and ih >= 20.0 and i_area >= 400.0:
                            figures.append(FigureData(
                                bbox=BBox(x0=ix0, y0=iy0, x1=ix1, y1=iy1),
                                page_num=page_num,
                                figure_type="raster",
                                figure_id=f"fig_p{page_num}_{len(figures) + 1}",
                                confidence=0.95
                            ))

            # Extract native PDF tables if available via pymupdf
            try:
                tabs = fitz_page.find_tables()
                for t_idx, tab in enumerate(tabs):
                    tab_df = tab.extract()
                    if tab_df and len(tab_df) > 0:
                        headers = [str(c or "").strip() for c in tab_df[0]]
                        rows = [[str(c or "").strip() for c in row] for row in tab_df[1:]]
                        tab_rect = tab.bbox
                        tx0 = max(0.0, min(page_w, float(tab_rect[0])))
                        ty0 = max(0.0, min(page_h, float(tab_rect[1])))
                        tx1 = max(0.0, min(page_w, float(tab_rect[2])))
                        ty1 = max(0.0, min(page_h, float(tab_rect[3])))
                        tables.append(TableData(
                            rows=rows,
                            headers=headers,
                            bbox=BBox(x0=tx0, y0=ty0, x1=tx1, y1=ty1),
                            page_num=page_num,
                            table_id=f"tbl_p{page_num}_{t_idx + 1}"
                        ))
            except Exception as e:
                logger.debug("Native PDF table extraction fallback for page %d: %s", page_num, e)

            # Native vector drawings detection with spatial clustering (diagrams, flowcharts, schematics)
            try:
                drawings = fitz_page.get_drawings()
                if drawings:
                    valid_rects = []
                    for d in drawings:
                        r = d.get("rect")
                        if not r:
                            continue
                        rx0 = max(0.0, min(page_w, float(r[0])))
                        ry0 = max(0.0, min(page_h, float(r[1])))
                        rx1 = max(0.0, min(page_w, float(r[2])))
                        ry1 = max(0.0, min(page_h, float(r[3])))
                        rw = rx1 - rx0
                        rh = ry1 - ry0
                        if rw < 1.0 or rh < 1.0:
                            continue
                        # Skip thin decorative separator rules (header/footer rules, dividing lines)
                        if rh <= 3.0 and rw >= page_w * 0.25:
                            continue
                        # Skip thin vertical column separator rules
                        if rw <= 3.0 and rh >= page_h * 0.25:
                            continue
                        # Skip full-page border boxes
                        if rw >= page_w * 0.90 and rh >= page_h * 0.90:
                            continue
                        # Skip table grid lines
                        is_table_grid = any(
                            t.bbox and (t.bbox.x0 - 5.0 <= rx0 and rx1 <= t.bbox.x1 + 5.0 and
                                        t.bbox.y0 - 5.0 <= ry0 and ry1 <= t.bbox.y1 + 5.0)
                            for t in tables
                        )
                        if is_table_grid:
                            continue
                        valid_rects.append((rx0, ry0, rx1, ry1))

                    if valid_rects:
                        # Connected components clustering of vector drawing primitives
                        max_gap = 50.0
                        n_r = len(valid_rects)
                        visited = [False] * n_r
                        clusters: List[Dict[str, Any]] = []

                        for i in range(n_r):
                            if visited[i]:
                                continue
                            component = [i]
                            visited[i] = True
                            queue = [i]
                            while queue:
                                curr = queue.pop(0)
                                cx0, cy0, cx1, cy1 = valid_rects[curr]
                                for j in range(n_r):
                                    if not visited[j]:
                                        jx0, jy0, jx1, jy1 = valid_rects[j]
                                        if not (jx1 < cx0 - max_gap or jx0 > cx1 + max_gap or
                                                jy1 < cy0 - max_gap or jy0 > cy1 + max_gap):
                                            visited[j] = True
                                            component.append(j)
                                            queue.append(j)

                            comp_x0 = min(valid_rects[idx][0] for idx in component)
                            comp_y0 = min(valid_rects[idx][1] for idx in component)
                            comp_x1 = max(valid_rects[idx][2] for idx in component)
                            comp_y1 = max(valid_rects[idx][3] for idx in component)
                            clusters.append({
                                "bbox": (comp_x0, comp_y0, comp_x1, comp_y1),
                                "count": len(component)
                            })

                        for c in clusters:
                            cx0, cy0, cx1, cy1 = c["bbox"]
                            cw = cx1 - cx0
                            ch = cy1 - cy0
                            c_area = cw * ch
                            if cw >= 50.0 and ch >= 35.0:
                                if c_area >= 0.85 * page_area:
                                    # Large engineering drawing or schematic
                                    if c["count"] >= 15 and total_native_chars < 150:
                                        figures.append(FigureData(
                                            bbox=BBox(x0=cx0, y0=cy0, x1=cx1, y1=cy1),
                                            page_num=page_num,
                                            figure_type="technical_drawing",
                                            figure_id=f"fig_p{page_num}_{len(figures) + 1}",
                                            confidence=0.90
                                        ))
                                else:
                                    figures.append(FigureData(
                                        bbox=BBox(x0=cx0, y0=cy0, x1=cx1, y1=cy1),
                                        page_num=page_num,
                                        figure_type="vector",
                                        figure_id=f"fig_p{page_num}_{len(figures) + 1}",
                                        confidence=0.90
                                    ))
            except Exception as d_err:
                logger.debug("Vector drawings extraction error on page %d: %s", page_num, d_err)

            # Deduplicate figures that heavily overlap
            if len(figures) > 1:
                deduped_figures: List[FigureData] = []
                for fig in figures:
                    if not fig.bbox:
                        continue
                    fb = fig.bbox
                    fig_area = fb.width * fb.height
                    duplicate = False
                    for existing in deduped_figures:
                        if not existing.bbox:
                            continue
                        eb = existing.bbox
                        ex_area = eb.width * eb.height
                        ix0 = max(fb.x0, eb.x0)
                        iy0 = max(fb.y0, eb.y0)
                        ix1 = min(fb.x1, eb.x1)
                        iy1 = min(fb.y1, eb.y1)
                        if ix1 > ix0 and iy1 > iy0:
                            inter_area = (ix1 - ix0) * (iy1 - iy0)
                            min_area = min(fig_area, ex_area)
                            if min_area > 0 and (inter_area / min_area > 0.80):
                                duplicate = True
                                break
                    if not duplicate:
                        deduped_figures.append(fig)
                figures = deduped_figures

            # Optional Deep CV analysis: only if deep_analysis is explicitly requested
            if deep_analysis:
                try:
                    pix = fitz_page.get_pixmap(dpi=150)
                    img_data = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
                    if pix.n >= 3:
                        gray = cv2.cvtColor(img_data, cv2.COLOR_BGR2GRAY if pix.n == 3 else cv2.COLOR_RGBA2GRAY)
                        cv_tables_boxes, _ = LayoutAnalyzerCV.detect_lines_and_tables(gray, page_num)
                        scale_x = page_w / pix.width
                        scale_y = page_h / pix.height
                        cv_figures = LayoutAnalyzerCV.detect_figures_and_drawings(gray, cv_tables_boxes, page_num)
                        for fig in cv_figures:
                            figures.append(FigureData(
                                bbox=BBox(
                                    x0=fig.bbox.x0 * scale_x, y0=fig.bbox.y0 * scale_y,
                                    x1=fig.bbox.x1 * scale_x, y1=fig.bbox.y1 * scale_y
                                ),
                                page_num=page_num,
                                figure_type=fig.figure_type,
                                figure_id=f"fig_p{page_num}_{len(figures) + 1}"
                            ))
                    del pix
                except Exception as cv_err:
                    logger.debug("Deep CV layout analysis on page %d skipped: %s", page_num, cv_err)

            # Reconstruct human reading order using geometric layout analysis
            ordered_blocks, reading_order, columns, layout_type = reading_engine.analyze_page_layout(
                blocks=raw_blocks,
                page_width=page_w,
                page_height=page_h,
                figures=figures,
                tables=tables
            )

            # Extract header and footer blocks based on geometric position
            header_y_max = page_h * 0.08
            footer_y_min = page_h * 0.92
            headers = [b for b in ordered_blocks if b.bbox.y1 <= header_y_max]
            footers = [b for b in ordered_blocks if b.bbox.y0 >= footer_y_min]

            page_text = "\n".join(b.text for b in ordered_blocks)
            is_blank_page = (
                len(raw_blocks) == 0 and
                len(tables) == 0 and
                len(figures) == 0 and
                not is_scanned_page
            )

            pages.append(PageModel(
                page_num=page_num,
                width=page_w,
                height=page_h,
                blocks=ordered_blocks,
                tables=tables,
                figures=figures,
                headers=headers,
                footers=footers,
                text=page_text,
                reading_order=reading_order,
                columns=columns,
                layout_type=layout_type,
                rotation=rotation,
                is_scanned=is_scanned_page,
                is_blank=is_blank_page
            ))

        doc.close()
        return DocumentModel(
            file_path=str(path),
            file_type="PDF",
            file_hash=file_hash,
            file_size=file_size,
            page_count=len(pages),
            pages=pages,
            metadata={"parser": "PyMuPDF-Incremental"}
        )

    @staticmethod
    def _parse_xlsx(path: Path, file_hash: str, file_size: int) -> DocumentModel:
        import openpyxl
        wb = openpyxl.load_workbook(str(path), data_only=True)
        pages: List[PageModel] = []

        for sheet_idx, sheet_name in enumerate(wb.sheetnames):
            ws = wb[sheet_name]
            all_rows = []
            for row in ws.iter_rows(values_only=True):
                # Convert None to ""
                all_rows.append([str(cell or "").strip() for cell in row])

            headers = all_rows[0] if all_rows else []
            data_rows = all_rows[1:] if len(all_rows) > 1 else []

            table = TableData(
                rows=data_rows,
                headers=headers,
                bbox=BBox(x0=50.0, y0=50.0, x1=740.0, y1=550.0),
                page_num=sheet_idx + 1,
                caption=f"Sheet: {sheet_name}"
            )

            # Build text representation
            text_lines = [f"=== Sheet: {sheet_name} ===", " | ".join(headers)]
            for r in data_rows[:100]:
                text_lines.append(" | ".join(r))

            blocks = [
                TextBlock(
                    text=line,
                    bbox=BBox(x0=50.0, y0=50.0 + idx * 16.0, x1=740.0, y1=64.0 + idx * 16.0),
                    font_name="Calibri",
                    font_size=10.0,
                    block_id=idx
                )
                for idx, line in enumerate(text_lines)
            ]

            pages.append(PageModel(
                page_num=sheet_idx + 1,
                width=792.0,
                height=612.0,
                blocks=blocks,
                tables=[table],
                text="\n".join(text_lines)
            ))

        wb.close()
        return DocumentModel(
            file_path=str(path),
            file_type="XLSX",
            file_hash=file_hash,
            file_size=file_size,
            page_count=len(pages),
            pages=pages,
            metadata={"parser": "openpyxl"}
        )

    @staticmethod
    def _parse_txt(path: Path, file_hash: str, file_size: int) -> DocumentModel:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            lines = [l.rstrip("\r\n") for l in f.readlines()]

        blocks = []
        page_lines = []
        pages = []
        page_num = 1
        y_pos = 50.0

        for idx, line in enumerate(lines):
            if idx > 0 and idx % 45 == 0:
                pages.append(PageModel(
                    page_num=page_num,
                    width=612.0,
                    height=792.0,
                    blocks=blocks,
                    tables=[],
                    text="\n".join(page_lines)
                ))
                page_num += 1
                blocks = []
                page_lines = []
                y_pos = 50.0

            if line.strip():
                blocks.append(TextBlock(
                    text=line,
                    bbox=BBox(x0=54.0, y0=y_pos, x1=558.0, y1=y_pos + 14.0),
                    font_name="Courier",
                    font_size=10.0,
                    block_id=idx
                ))
            page_lines.append(line)
            y_pos += 15.0

        if page_lines or not pages:
            pages.append(PageModel(
                page_num=page_num,
                width=612.0,
                height=792.0,
                blocks=blocks,
                tables=[],
                text="\n".join(page_lines)
            ))

        return DocumentModel(
            file_path=str(path),
            file_type="TXT",
            file_hash=file_hash,
            file_size=file_size,
            page_count=len(pages),
            pages=pages,
            metadata={"parser": "plain_text"}
        )

    @staticmethod
    def _parse_image(path: Path, file_hash: str, file_size: int) -> DocumentModel:
        image = cv2.imread(str(path))
        if image is None:
            raise ValueError(f"Could not load image at {path}")

        h, w = image.shape[:2]
        # Preprocess, deskew, and binarize
        enhanced = ScannedDocumentPipeline.preprocess_image(image)
        deskewed, angle = ScannedDocumentPipeline.deskew_image(enhanced)
        binary = ScannedDocumentPipeline.binarize(deskewed)

        # Extract text regions
        text_boxes = ScannedDocumentPipeline.extract_text_regions(binary)
        table_boxes, tables = LayoutAnalyzerCV.detect_lines_and_tables(deskewed, page_num=1)
        figures = LayoutAnalyzerCV.detect_figures_and_drawings(deskewed, table_boxes, page_num=1)

        # Create blocks from text regions
        blocks = [
            TextBlock(
                text=f"[Image Text Region {idx+1}]",
                bbox=bbox,
                font_name="OCR-Detected",
                font_size=11.0,
                block_id=idx
            )
            for idx, bbox in enumerate(text_boxes)
        ]

        page = PageModel(
            page_num=1,
            width=float(w),
            height=float(h),
            blocks=blocks,
            tables=tables,
            figures=figures,
            text=f"[Scanned Image Document: {path.name} | Deskew: {angle:.1f}° | Regions: {len(text_boxes)}]"
        )

        return DocumentModel(
            file_path=str(path),
            file_type="IMAGE",
            file_hash=file_hash,
            file_size=file_size,
            page_count=1,
            pages=[page],
            metadata={"parser": "OpenCV-Scanned", "deskew_angle": angle}
        )
